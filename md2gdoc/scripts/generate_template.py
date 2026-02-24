"""
Auto-generate a template skeleton from a .docx file.

Parses the document's heading structure and metadata fields,
then generates a skeleton.md with {{PLACEHOLDER}} tokens and
caches the .docx as a pandoc reference document.

Usage:
    python3 generate_template.py <source> <template_name> [--download] [--templates-dir DIR]

    source:         local .docx path or Google Drive file ID (with --download)
    template_name:  name for the cached template directory
    --download:     treat source as a Google Drive file ID and download via Drive connector
    --templates-dir: override the default templates directory
"""

import argparse
import asyncio
import re
import shutil
import sys
from pathlib import Path

from docx import Document


DEFAULT_TEMPLATES_DIR = Path.home() / ".claude" / "skills" / "md2gdoc" / "templates"


def to_upper_snake_case(text: str) -> str:
    """Convert heading text to UPPER_SNAKE_CASE placeholder name.

    'Data Protection Considerations' -> 'DATA_PROTECTION_CONSIDERATIONS'
    'Open questions' -> 'OPEN_QUESTIONS'
    'Reviewers / Informed' -> 'REVIEWERS_INFORMED'
    """
    cleaned = re.sub(r"[^a-zA-Z0-9\s]", " ", text)
    parts = cleaned.strip().split()
    # Limit to first 6 words to avoid overly long placeholder names
    parts = parts[:6]
    return "_".join(p.upper() for p in parts if p)


def is_metadata_line(text: str, is_docx: bool = False) -> bool:
    """Detect lines that look like metadata key-value fields.

    Supports two formats:
    - Markdown: '**Key:** Value'
    - Plain text (from .docx): 'Key:' or 'Key: Value'

    Only valid for content before the first heading in the document.
    """
    # Markdown bold key pattern
    if re.match(r"\*\*[^*]+:\*\*", text):
        return True
    # Plain text key pattern (from .docx paragraphs)
    if is_docx and re.match(r"^[A-Z][a-zA-Z\s]+:\s*", text):
        return True
    return False


def extract_key_from_text(text: str, is_docx: bool = False) -> tuple:
    """Extract the key pattern and placeholder name from a metadata line.

    Markdown: '**Author:** John Doe' -> ('**Author:**', 'AUTHOR')
    Plain:    'Author: '             -> ('**Author:**', 'AUTHOR')

    Returns (display_pattern, placeholder_name) or (None, None) if not a match.
    """
    # Try markdown bold pattern first
    m = re.match(r"(\*\*([^*]+):\*\*)", text)
    if m:
        display = m.group(1)
        key_text = m.group(2).strip()
        placeholder = to_upper_snake_case(key_text)
        return display, placeholder

    # Try plain text pattern (from .docx)
    if is_docx:
        m = re.match(r"^([A-Z][a-zA-Z\s]+):\s*", text)
        if m:
            key_text = m.group(1).strip()
            placeholder = to_upper_snake_case(key_text)
            # Output as bold markdown for the skeleton
            display = f"**{key_text}:**"
            return display, placeholder

    return None, None


def extract_headings_and_metadata(docx_path: str) -> dict:
    """Parse a .docx file and extract its structure.

    Returns:
        {
            'metadata': [{'display': '**Author:**', 'placeholder': 'AUTHOR'}, ...],
            'headings': [{'level': 1, 'text': 'Summary', 'placeholder': 'SUMMARY'}, ...],
        }
    """
    doc = Document(docx_path)

    metadata = []
    headings = []
    seen_heading = False
    seen_placeholders = set()

    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.strip()

        if not text:
            continue

        # Check for headings
        if style_name.startswith("Heading "):
            seen_heading = True
            try:
                level = int(style_name.split()[-1])
            except ValueError:
                continue

            placeholder = to_upper_snake_case(text)
            if not placeholder:
                continue

            # Handle duplicates with numeric suffix
            original = placeholder
            counter = 2
            while placeholder in seen_placeholders:
                placeholder = f"{original}_{counter}"
                counter += 1
            seen_placeholders.add(placeholder)

            headings.append({
                "level": level,
                "text": text,
                "placeholder": placeholder,
            })
            continue

        # Check for metadata lines (only before first heading)
        if not seen_heading and is_metadata_line(text, is_docx=True):
            display, placeholder = extract_key_from_text(text, is_docx=True)
            if display and placeholder:
                # Handle duplicate metadata keys
                original = placeholder
                counter = 2
                while placeholder in seen_placeholders:
                    placeholder = f"{original}_{counter}"
                    counter += 1
                seen_placeholders.add(placeholder)

                metadata.append({
                    "display": display,
                    "placeholder": placeholder,
                })

    return {"metadata": metadata, "headings": headings}


def generate_skeleton(structure: dict) -> str:
    """Generate a template skeleton markdown from the parsed structure."""
    lines = []

    # YAML frontmatter with title placeholder
    lines.append("---")
    lines.append('title: "{{TITLE}}"')
    lines.append("---")
    lines.append("")

    # Metadata section
    for meta in structure["metadata"]:
        lines.append(f'{meta["display"]} {{{{{meta["placeholder"]}}}}}')
        lines.append("")

    # Heading sections
    for heading in structure["headings"]:
        prefix = "#" * heading["level"]
        lines.append(f'{prefix} {heading["text"]}')
        lines.append("")
        lines.append(f'{{{{{heading["placeholder"]}}}}}')
        lines.append("")

    return "\n".join(lines)


def download_from_gdrive(file_id: str, output_path: str) -> bool:
    """Download a Google Doc as .docx via the Drive connector.

    Uses the airbyte-agent-google-drive connector with .env credentials
    to export a Google Doc as .docx format.
    """
    from gdrive_auth import get_connector

    try:
        connector = get_connector()
    except Exception as e:
        print(f"Error initializing Drive connector: {e}", file=sys.stderr)
        return False

    async def _download():
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        await connector.files_export.download_local(
            file_id=file_id,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            path=output_path,
        )

    try:
        asyncio.run(_download())
        return True
    except Exception as e:
        print(f"Error downloading from Google Drive: {e}", file=sys.stderr)
        return False


def main():
    parser = argparse.ArgumentParser(
        description="Auto-generate a template skeleton from a .docx file"
    )
    parser.add_argument("source", help="Local .docx path or Google Drive file ID (with --download)")
    parser.add_argument("template_name", help="Name for the cached template")
    parser.add_argument("--download", action="store_true",
                        help="Treat source as Google Drive file ID and download via Drive connector")
    parser.add_argument("--templates-dir", type=Path, default=DEFAULT_TEMPLATES_DIR,
                        help="Override the templates directory")
    args = parser.parse_args()

    # Resolve source .docx
    if args.download:
        tmp_path = f"/tmp/md2gdoc/template_source_{args.template_name}.docx"
        Path(tmp_path).parent.mkdir(parents=True, exist_ok=True)
        print(f"Downloading Google Doc {args.source} as .docx...")
        if not download_from_gdrive(args.source, tmp_path):
            sys.exit(1)
        source_path = tmp_path
        print(f"  Downloaded to: {source_path}")
    else:
        source_path = args.source
        if not Path(source_path).exists():
            print(f"Error: file not found: {source_path}", file=sys.stderr)
            sys.exit(1)

    # Parse document structure
    print(f"Parsing: {source_path}")
    structure = extract_headings_and_metadata(source_path)

    print(f"  Found {len(structure['metadata'])} metadata fields")
    print(f"  Found {len(structure['headings'])} headings")

    if not structure["headings"]:
        print("Warning: no headings found in the document. The skeleton will only have metadata.", file=sys.stderr)

    # Generate skeleton
    skeleton_text = generate_skeleton(structure)

    # Save to cache
    cache_dir = args.templates_dir / args.template_name
    cache_dir.mkdir(parents=True, exist_ok=True)

    skeleton_path = cache_dir / "skeleton.md"
    skeleton_path.write_text(skeleton_text, encoding="utf-8")
    print(f"  Skeleton saved to: {skeleton_path}")

    reference_path = cache_dir / "reference.docx"
    shutil.copy2(source_path, reference_path)
    print(f"  Reference .docx saved to: {reference_path}")

    # Print summary
    print(f"\nTemplate '{args.template_name}' cached at: {cache_dir}")
    print(f"\nPlaceholders generated:")
    for meta in structure["metadata"]:
        print(f"  [metadata] {meta['display']} → {{{{{meta['placeholder']}}}}}")
    for heading in structure["headings"]:
        prefix = "#" * heading["level"]
        print(f"  [{prefix}] {heading['text']} → {{{{{heading['placeholder']}}}}}")

    # Clean up temp file
    if args.download:
        Path(source_path).unlink(missing_ok=True)


if __name__ == "__main__":
    main()
